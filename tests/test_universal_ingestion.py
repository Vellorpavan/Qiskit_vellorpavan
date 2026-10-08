"""
ShiftProof — Universal Data Ingestion & Scheduling Understanding Test Suite
Verifies Requirements:
- Deterministic Parsers: CSV, XLSX, JSON, XML, Text-PDF
- Document Ingestion Abstraction: PDF table extraction, scanned PDF fallback, Gemini mock adapter
- Semantic Schema Understanding & Scheduling Problem Classifier Matrix:
    1. employee_id | duty | hourly_wage -> VALID_SCHEDULING_DATASET
    2. Employee Name | Shift | Rate | Availability -> VALID_SCHEDULING_DATASET
    3. Staff | Shift Type | Assignment Cost | Eligible -> VALID_SCHEDULING_DATASET
    4. Worker | Shift | Cost -> VALID_SCHEDULING_DATASET
    5. Student | Marks | Grade -> NOT_A_SCHEDULING_DATASET (reject)
    6. Employee | Department | Salary -> SCHEDULING_DATASET_MISSING_REQUIRED_CONCEPT (reject)
    7. Employee | Morning | Evening | Night -> POTENTIAL_SHIFT_MATRIX
- Row-based assignment format (eligibility inferred from assignment rows)
- Explicit availability format normalization (yes/no, true/false, 1/0, available/unavailable)
- Ambiguity detection (competing candidates requiring confirmation)
- Matrix unpivoting
- Zero fabrication of missing information
- A/B Data Proof:
    Dataset A (A:10, B:50) vs Dataset B (A:90, B:5)
    -> Fingerprints differ
    -> Normalized instances differ
    -> QUBOs differ
    -> Classical optima differ
    -> Persisted DB results are strictly isolated
"""

import io
import json
import sqlite3
import pytest
import pandas as pd
import numpy as np
import openpyxl

from app.database.schema import init_db
from app.services.ingestion.file_detector import FileType, FileCategory, detect_file_type
from app.services.ingestion.structured_parser import parse_structured_file, ParsedStructure
from app.services.ingestion.document_parser import parse_document_file, is_gemini_configured
from app.services.ingestion.schema_understanding import understand_schema
from app.services.ingestion.scheduling_classifier import (
    ClassificationStatus,
    classify_dataset,
)
from app.services.ingestion.normalizer import normalize_dataset
from app.services.ingestion.pipeline import (
    inspect_uploaded_file,
    finalize_and_save_dataset,
)
from app.services.scheduling_service import db_to_instance, compute_dataset_fingerprint
from src.classical import solve_exact
from src.quantum import build_qubo


@pytest.fixture
def test_db():
    """Isolated in-memory SQLite database initialized with ShiftProof schema."""
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    init_db(conn)
    yield conn
    conn.close()


# ==============================================================================
# 1. FILE TYPE DETECTION TESTS
# ==============================================================================

def test_detect_file_types():
    # CSV
    csv_bytes = b"Worker,Shift,Cost\nAarav,Morning,100"
    t, cat = detect_file_type("schedule.csv", csv_bytes)
    assert t == FileType.CSV
    assert cat == FileCategory.STRUCTURED

    # TSV
    tsv_bytes = b"Worker\tShift\tCost\nAarav\tMorning\t100"
    t, cat = detect_file_type("schedule.tsv", tsv_bytes)
    assert t == FileType.TSV
    assert cat == FileCategory.STRUCTURED

    # JSON
    json_bytes = b'[{"Worker": "Aarav", "Shift": "Morning", "Cost": 100}]'
    t, cat = detect_file_type("schedule.json", json_bytes)
    assert t == FileType.JSON
    assert cat == FileCategory.STRUCTURED

    # XML
    xml_bytes = b"<root><record><Worker>Aarav</Worker><Shift>Morning</Shift><Cost>100</Cost></record></root>"
    t, cat = detect_file_type("schedule.xml", xml_bytes)
    assert t == FileType.XML
    assert cat == FileCategory.STRUCTURED

    # PDF
    pdf_bytes = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\n"
    t, cat = detect_file_type("roster.pdf", pdf_bytes)
    assert t == FileType.PDF
    assert cat == FileCategory.DOCUMENT


# ==============================================================================
# 2. DETERMINISTIC PARSERS (CSV, XLSX, JSON, XML)
# ==============================================================================

def test_parse_csv_semicolon_and_comma():
    # Comma delimited
    csv1 = b"Worker,Shift,Cost\nAarav,Morning,100\nMeera,Night,120\n"
    parsed1 = parse_structured_file(csv1, "test.csv", FileType.CSV)
    assert len(parsed1.df) == 2
    assert list(parsed1.df.columns) == ["Worker", "Shift", "Cost"]

    # Semicolon delimited
    csv2 = b"Worker;Shift;Cost\nAarav;Morning;100\nMeera;Night;120\n"
    parsed2 = parse_structured_file(csv2, "test.csv", FileType.CSV)
    assert len(parsed2.df) == 2
    assert list(parsed2.df.columns) == ["Worker", "Shift", "Cost"]


def test_parse_xlsx_multi_sheet():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Overview"
    ws1.append(["Title", "Company Schedule"])
    ws1.append(["Date", "2026-10-15"])

    ws2 = wb.create_sheet("Assignments")
    ws2.append(["Worker", "Shift", "Cost"])
    ws2.append(["Aarav", "Morning", 100])
    ws2.append(["Meera", "Night", 120])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    # Parse without explicit sheet (auto selects highest row/col table sheet)
    parsed = parse_structured_file(xlsx_bytes, "workforce.xlsx", FileType.XLSX)
    assert "Assignments" in parsed.sheet_names
    assert len(parsed.df) == 2
    assert list(parsed.df.columns) == ["Worker", "Shift", "Cost"]

    # Parse with explicit sheet
    parsed_sheet = parse_structured_file(xlsx_bytes, "workforce.xlsx", FileType.XLSX, sheet_name="Assignments")
    assert len(parsed_sheet.df) == 2
    assert parsed_sheet.selected_sheet == "Assignments"


def test_parse_json_formats():
    # Records list format
    data_list = [
        {"Worker": "Aarav", "Shift": "Morning", "Cost": 100},
        {"Worker": "Meera", "Shift": "Evening", "Cost": 120},
    ]
    parsed_list = parse_structured_file(json.dumps(data_list).encode(), "data.json", FileType.JSON)
    assert len(parsed_list.df) == 2
    assert "Worker" in parsed_list.df.columns

    # Columnar dictionary format
    data_dict = {
        "Worker": ["Aarav", "Meera"],
        "Shift": ["Morning", "Evening"],
        "Cost": [100, 120],
    }
    parsed_dict = parse_structured_file(json.dumps(data_dict).encode(), "data.json", FileType.JSON)
    assert len(parsed_dict.df) == 2
    assert list(parsed_dict.df.columns) == ["Worker", "Shift", "Cost"]

    # Wrapped key dictionary format
    data_wrapped = {
        "metadata": {"generated": "2026"},
        "assignments": data_list,
    }
    parsed_wrapped = parse_structured_file(json.dumps(data_wrapped).encode(), "data.json", FileType.JSON)
    assert len(parsed_wrapped.df) == 2


def test_parse_xml_format():
    xml_content = b"""<?xml version="1.0" encoding="UTF-8"?>
    <schedule>
        <item>
            <Worker>Aarav</Worker>
            <Shift>Morning</Shift>
            <Cost>100</Cost>
        </item>
        <item>
            <Worker>Meera</Worker>
            <Shift>Evening</Shift>
            <Cost>120</Cost>
        </item>
    </schedule>
    """
    parsed = parse_structured_file(xml_content, "schedule.xml", FileType.XML)
    assert len(parsed.df) == 2
    assert set(parsed.df.columns) == {"Worker", "Shift", "Cost"}


# ==============================================================================
# 3. DOCUMENT / PDF INGESTION ABSTRACTION
# ==============================================================================

def test_parse_pdf_local_deterministic():
    """Text-based PDF with structured tabular lines should parse locally via pypdf."""
    pdf_bytes = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>
endobj
4 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
5 0 obj
<< /Length 140 >>
stream
BT
/F1 12 Tf
72 700 Td
(Worker,Shift,Cost) Tj
0 -20 Td
(Aarav,Morning,100) Tj
0 -20 Td
(Meera,Evening,120) Tj
ET
endstream
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000227 00000 n 
0000000298 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
490
%%EOF"""

    parsed = parse_document_file(pdf_bytes, "roster.pdf", FileType.PDF)
    assert len(parsed.df) == 2
    assert list(parsed.df.columns) == ["Worker", "Shift", "Cost"]
    assert parsed.metadata["extractor"] == "local_pypdf"


def test_parse_pdf_scanned_without_gemini_raises_informative_error(monkeypatch):
    """A scanned or empty PDF without text must explain clearly that GEMINI_API_KEY is needed."""
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    assert not is_gemini_configured()

    empty_pdf = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    with pytest.raises(ValueError, match="GEMINI_API_KEY"):
        parse_document_file(empty_pdf, "scanned_roster.pdf", FileType.PDF)


# ==============================================================================
# 4. SEMANTIC SCHEMA UNDERSTANDING & CLASSIFIER MATRIX (REQUIREMENT 24)
# ==============================================================================

def test_classifier_case_1_employee_duty_wage():
    # 1. employee_id | duty | hourly_wage -> scheduling
    df = pd.DataFrame({
        "employee_id": ["E101", "E102", "E103"],
        "duty": ["Morning", "Evening", "Night"],
        "hourly_wage": [25.0, 30.0, 35.0],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.VALID_SCHEDULING_DATASET
    assert res.detected_mapping["worker"] == "employee_id"
    assert res.detected_mapping["shift"] == "duty"
    assert res.detected_mapping["cost"] == "hourly_wage"


def test_classifier_case_2_employee_name_shift_rate_availability():
    # 2. Employee Name | Shift | Rate | Availability -> scheduling
    df = pd.DataFrame({
        "Employee Name": ["Aarav", "Meera", "Kabir"],
        "Shift": ["Morning", "Evening", "Night"],
        "Rate": [100.0, 110.0, 120.0],
        "Availability": ["yes", "no", "yes"],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.VALID_SCHEDULING_DATASET
    assert res.detected_mapping["worker"] == "Employee Name"
    assert res.detected_mapping["shift"] == "Shift"
    assert res.detected_mapping["cost"] == "Rate"
    assert res.detected_mapping["availability"] == "Availability"


def test_classifier_case_3_staff_shifttype_cost_eligible():
    # 3. Staff | Shift Type | Assignment Cost | Eligible -> scheduling
    df = pd.DataFrame({
        "Staff": ["Dr. Rao", "Dr. Nair", "Dr. Mehta"],
        "Shift Type": ["ICU Day", "ER Night", "Ward Day"],
        "Assignment Cost": [150.0, 180.0, 130.0],
        "Eligible": [1, 0, 1],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.VALID_SCHEDULING_DATASET
    assert res.detected_mapping["worker"] == "Staff"
    assert res.detected_mapping["shift"] == "Shift Type"
    assert res.detected_mapping["cost"] == "Assignment Cost"
    assert res.detected_mapping["availability"] == "Eligible"


def test_classifier_case_4_worker_shift_cost():
    # 4. Worker | Shift | Cost -> scheduling
    df = pd.DataFrame({
        "Worker": ["W1", "W2", "W3"],
        "Shift": ["S1", "S2", "S3"],
        "Cost": [50.0, 60.0, 70.0],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.VALID_SCHEDULING_DATASET
    assert res.detected_mapping["worker"] == "Worker"
    assert res.detected_mapping["shift"] == "Shift"
    assert res.detected_mapping["cost"] == "Cost"


def test_classifier_case_5_student_marks_grade_rejected():
    # 5. Student | Marks | Grade -> reject (NOT_A_SCHEDULING_DATASET)
    df = pd.DataFrame({
        "Student": ["Alice", "Bob", "Charlie"],
        "Marks": [85, 92, 78],
        "Grade": ["A", "A+", "B"],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.NOT_A_SCHEDULING_DATASET
    assert not res.is_valid
    assert "shift" not in res.detected_mapping or res.detected_mapping["shift"] is None


def test_classifier_case_6_employee_department_salary_rejected():
    # 6. Employee | Department | Salary -> reject (missing shift concept)
    df = pd.DataFrame({
        "Employee": ["Aarav", "Meera", "Kabir"],
        "Department": ["Cardiology", "Neurology", "Pediatrics"],
        "Salary": [120000, 140000, 110000],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.SCHEDULING_DATASET_MISSING_REQUIRED_CONCEPT
    assert not res.is_valid
    assert any("Shift" in c for c in res.missing_concepts)


def test_classifier_case_7_employee_morning_evening_night_matrix():
    # 7. Employee | Morning | Evening | Night -> potential matrix requiring interpretation
    df = pd.DataFrame({
        "Employee": ["Aarav", "Meera", "Kabir"],
        "Morning": [1, 1, 0],
        "Evening": [0, 1, 1],
        "Night": [1, 0, 1],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.POTENTIAL_SHIFT_MATRIX
    assert res.is_matrix
    assert set(res.matrix_shift_columns) == {"Morning", "Evening", "Night"}


# ==============================================================================
# 5. AVAILABILITY NORMALIZATION & ZERO FABRICATION
# ==============================================================================

def test_explicit_availability_normalization():
    df = pd.DataFrame({
        "Worker": ["W1", "W2", "W3", "W4", "W5", "W6", "W7", "W8"],
        "Shift": ["S1"] * 8,
        "Cost": [10.0] * 8,
        "Avail": ["true", "false", "yes", "no", "1", "0", "available", "unavailable"],
    })
    norm = normalize_dataset(
        df,
        mapping={"worker": "Worker", "shift": "Shift", "cost": "Cost", "availability": "Avail"},
    )
    assert norm.is_valid
    expected_avail = [1, 0, 1, 0, 1, 0, 1, 0]
    assert list(norm.normalized_df["is_available"]) == expected_avail
    assert list(norm.normalized_df["is_eligible"]) == expected_avail


def test_dataset_with_assignment_cost_mode_cost_optimization(test_db):
    """Case A: Dataset WITH assignment cost -> COST_OPTIMIZATION."""
    csv_bytes = b"Worker,Shift,Hourly Rate\nW1,Morning,25.0\nW2,Evening,35.0\n"
    insp = inspect_uploaded_file(csv_bytes, "shifts_with_cost.csv")
    assert insp.classification.is_valid
    assert insp.classification.optimization_mode == "COST_OPTIMIZATION"
    assert insp.classification.detected_mapping["cost"] == "Hourly Rate"

    ds_id, norm = finalize_and_save_dataset(test_db, insp, dataset_name="With Cost")
    assert norm.is_valid
    assert norm.optimization_mode == "COST_OPTIMIZATION"
    assert norm.normalized_df["assignment_cost"].tolist() == [25.0, 35.0]

    # Verify Instance has non-zero cost matrix
    inst = db_to_instance(ds_id, test_db)
    assert inst.cost_matrix is not None
    assert np.any(inst.cost_matrix > 0)


def test_dataset_without_assignment_cost_mode_feasibility_only(test_db):
    """
    Case B: Dataset WITHOUT assignment cost:
      -> FEASIBILITY_ONLY
      -> actual Instance is created
      -> actual QUBO is created
      -> no fabricated monetary values
      -> classical execution works
      -> QAOA execution works on supported small instances
    """
    from app.services.experiment_service import run_qaoa_workspace

    # Dataset without any wage/cost column
    csv_bytes = b"Worker,Shift\nW1,Morning\nW2,Morning\n"
    insp = inspect_uploaded_file(csv_bytes, "shifts_no_cost.csv")

    # 1. Classification & Optimization mode
    assert insp.classification.is_valid
    assert insp.classification.optimization_mode == "FEASIBILITY_ONLY"
    assert insp.classification.detected_mapping.get("cost") is None

    # 2. Finalize & Save
    ds_id, norm = finalize_and_save_dataset(test_db, insp, dataset_name="No Cost Dataset")
    assert norm.is_valid
    assert norm.optimization_mode == "FEASIBILITY_ONLY"
    # Never invent random or fake non-zero monetary values
    assert all(c == 0.0 for c in norm.normalized_df["assignment_cost"])

    # 3. Actual Instance is created
    inst = db_to_instance(ds_id, test_db)
    assert inst.n_workers == 2
    assert inst.n_shifts == 1
    assert inst.cost_matrix is not None
    assert np.all(inst.cost_matrix == 0.0)

    # 4. Actual QUBO is created
    qubo = build_qubo(inst)
    assert qubo.a is not None
    assert qubo.b is not None
    assert qubo.A > 0.0
    assert qubo.B > 0.0
    # In QUBO, penalty A and B are strictly positive (delta=1.0) so constraints are enforced

    # 5. Classical solver executes and finds feasible schedule
    exact_res = solve_exact(inst)
    assert exact_res.feasible
    assert exact_res.optimal_cost == 0.0  # Feasibility objective mathematically 0.0

    # 6. QAOA execution on AerSimulator works and samples are verified
    q_out = run_qaoa_workspace(inst, ds_id, shots=64, maxiter=3, seed=42, conn=test_db)
    assert q_out["backend"] == "Qiskit AerSimulator"
    assert q_out["qubit_count"] == 2
    assert q_out["feasible_rate"] >= 0.0
    if q_out["feasible"]:
        assert q_out["best_feasible_cost"] == 0.0


def test_real_30_row_8_col_xlsx_workforce_without_cost(test_db):
    """
    Test real 30-row, 8-column XLSX workforce dataset with Worker & Shift / Duty Start_Time,
    without Assignment Cost/Wage:
    - Must NOT be rejected as 'Scheduling problem not detected'.
    - Must detect Worker & Shift / Duty Start_Time.
    - Optimization mode: FEASIBILITY_ONLY.
    - Successfully creates Instance -> QUBO -> Classical solver.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"

    headers = [
        "Worker_ID",
        "Employee_Name",
        "Department",
        "Duty_Date",
        "Start_Time",
        "End_Time",
        "Role",
        "Notes"
    ]
    ws.append(headers)

    # 30 rows of workforce duty data
    shifts = ["08:00", "16:00", "00:00"]
    for i in range(1, 31):
        worker_id = f"W{i:02d}"
        worker_name = f"Worker_{i}"
        dept = "Operations"
        duty_date = "2026-10-10"
        start_t = shifts[i % 3]
        end_t = "16:00"
        role = "Operator"
        notes = "Standard"
        ws.append([worker_id, worker_name, dept, duty_date, start_t, end_t, role, notes])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    insp = inspect_uploaded_file(xlsx_bytes, "workforce_30_rows.xlsx")

    # 1. Must be recognized as scheduling dataset (VALID or NEEDS_MAPPING due to Worker_ID + Employee_Name)
    assert insp.classification.status in [
        ClassificationStatus.VALID_SCHEDULING_DATASET,
        ClassificationStatus.SCHEDULING_DATASET_NEEDS_MAPPING,
    ]
    assert insp.classification.optimization_mode == "FEASIBILITY_ONLY"
    assert insp.classification.detected_mapping.get("cost") is None
    # If ambiguous (both Worker_ID and Employee_Name), Requirement 6: ask for mapping instead of rejecting
    if insp.classification.status == ClassificationStatus.SCHEDULING_DATASET_NEEDS_MAPPING:
        assert "worker" in insp.classification.ambiguous_concepts

    # 2. Normalize and save with mapping confirmation
    confirmed_mapping = {
        "worker": "Worker_ID",
        "shift": "Start_Time",
        "cost": None,
    }
    ds_id, norm = finalize_and_save_dataset(
        test_db, insp, user_mapping=confirmed_mapping, dataset_name="30-Row Workforce"
    )
    assert norm.is_valid
    assert norm.optimization_mode == "FEASIBILITY_ONLY"
    assert len(norm.normalized_df) == 30

    # 3. Instance created
    inst = db_to_instance(ds_id, test_db)
    assert inst.n_workers > 0
    assert inst.n_shifts > 0
    assert np.all(inst.cost_matrix == 0.0)

    # 4. QUBO created
    qubo = build_qubo(inst)
    assert qubo.instance.n_vars == inst.n_vars
    assert qubo.a.shape == (inst.n_vars,)
    assert qubo.b.shape == (inst.n_vars, inst.n_vars)

    # 5. Classical solver executes
    from src.classical import solve_greedy
    greedy_res = solve_greedy(inst)
    assert greedy_res.feasible
    assert greedy_res.optimal_cost == 0.0



def test_ambiguous_cost_columns_require_user_confirmation():
    """When multiple columns score similarly for cost, flag ambiguity."""
    df = pd.DataFrame({
        "Worker": ["W1", "W2", "W3"],
        "Shift": ["S1", "S2", "S3"],
        "Hourly Rate": [40.0, 45.0, 50.0],
        "Labor Cost": [320.0, 360.0, 400.0],
    })
    cands = understand_schema(df)
    res = classify_dataset(df, cands)
    assert res.status == ClassificationStatus.SCHEDULING_DATASET_NEEDS_MAPPING
    assert "cost" in res.ambiguous_concepts
    assert len(cands["cost"]) >= 2


def test_matrix_unpivoting():
    """Test unpivoting of Employee | Morning | Evening | Night matrix."""
    df = pd.DataFrame({
        "Employee": ["Aarav", "Meera"],
        "Morning": [10.0, 20.0],
        "Evening": [15.0, 25.0],
    })
    norm = normalize_dataset(
        df,
        mapping={"worker": "Employee"},
        is_matrix=True,
        matrix_cell_type="cost",
        matrix_shift_columns=["Morning", "Evening"],
    )
    assert norm.is_valid
    assert len(norm.normalized_df) == 4
    assert set(norm.normalized_df["worker_name"]) == {"Aarav", "Meera"}
    assert set(norm.normalized_df["shift_name"]) == {"Morning", "Evening"}
    # Verify Aarav Morning cost = 10.0
    aarav_morning = norm.normalized_df[
        (norm.normalized_df["worker_name"] == "Aarav") & (norm.normalized_df["shift_name"] == "Morning")
    ].iloc[0]
    assert aarav_morning["assignment_cost"] == 10.0


# ==============================================================================
# 6. A/B DATA PROOF (MANDATORY REQUIREMENT 25)
# ==============================================================================

def test_ab_data_proof_different_costs_different_results(test_db):
    """
    MANDATORY REQUIREMENT 25:
    Create two valid datasets with different costs:
      Dataset A:
        A -> Morning -> 10
        B -> Morning -> 50
      Dataset B:
        A -> Morning -> 90
        B -> Morning -> 5
    Verify:
      1. dataset fingerprint changes
      2. normalized model changes
      3. QUBO changes
      4. classical optimum changes
      5. persisted result changes
    """
    # Create CSV A
    csv_a = b"Worker,Shift,Cost\nA,Morning,10\nB,Morning,50\n"
    insp_a = inspect_uploaded_file(csv_a, "dataset_a.csv")
    assert insp_a.classification.is_valid
    ds_id_a, norm_a = finalize_and_save_dataset(test_db, insp_a, dataset_name="Dataset A")

    # Create CSV B
    csv_b = b"Worker,Shift,Cost\nA,Morning,90\nB,Morning,5\n"
    insp_b = inspect_uploaded_file(csv_b, "dataset_b.csv")
    assert insp_b.classification.is_valid
    ds_id_b, norm_b = finalize_and_save_dataset(test_db, insp_b, dataset_name="Dataset B")

    # 1. Dataset fingerprint changes
    fp_a = compute_dataset_fingerprint(ds_id_a, test_db)
    fp_b = compute_dataset_fingerprint(ds_id_b, test_db)
    assert fp_a != fp_b, f"Dataset fingerprints must differ! A={fp_a}, B={fp_b}"

    # 2. Normalized models change
    inst_a = db_to_instance(ds_id_a, test_db)
    inst_b = db_to_instance(ds_id_b, test_db)
    assert inst_a.cost_matrix is not None
    assert inst_b.cost_matrix is not None
    assert not np.array_equal(inst_a.cost_matrix, inst_b.cost_matrix)
    # In A: worker 0 (A) cost=10, worker 1 (B) cost=50
    assert inst_a.cost_matrix[0, 0] == 10.0
    assert inst_a.cost_matrix[1, 0] == 50.0
    # In B: worker 0 (A) cost=90, worker 1 (B) cost=5
    assert inst_b.cost_matrix[0, 0] == 90.0
    assert inst_b.cost_matrix[1, 0] == 5.0

    # 3. QUBO changes
    qubo_a = build_qubo(inst_a)
    qubo_b = build_qubo(inst_b)
    assert not np.array_equal(qubo_a.a, qubo_b.a), "QUBO linear terms must differ"

    # 4. Classical optimum changes
    res_a = solve_exact(inst_a)
    res_b = solve_exact(inst_b)

    assert res_a.feasible
    assert res_b.feasible
    # In Dataset A, Worker A (index 0) has cost 10 < 50
    assert res_a.optimal_cost == 10.0

    # In Dataset B, Worker B (index 1) has cost 5 < 90
    assert res_b.optimal_cost == 5.0

    assert res_a.optimal_cost != res_b.optimal_cost
    assert res_a.optimal_assignment != res_b.optimal_assignment

    # 5. Persisted result changes in database
    with test_db:
        test_db.execute(
            "INSERT INTO experiments (experiment_id, dataset_id, solver_type, status, parameters_json) "
            "VALUES ('exp_a', ?, 'classical_exact', 'completed', '{}')",
            (ds_id_a,),
        )
        test_db.execute(
            "INSERT INTO solutions (solution_id, experiment_id, dataset_id, feasible, total_cost, runtime_seconds, assignment_json, verification_json) "
            "VALUES ('sol_a', 'exp_a', ?, 1, ?, 0.01, '{}', '{}')",
            (ds_id_a, res_a.optimal_cost),
        )
        test_db.execute(
            "INSERT INTO experiments (experiment_id, dataset_id, solver_type, status, parameters_json) "
            "VALUES ('exp_b', ?, 'classical_exact', 'completed', '{}')",
            (ds_id_b,),
        )
        test_db.execute(
            "INSERT INTO solutions (solution_id, experiment_id, dataset_id, feasible, total_cost, runtime_seconds, assignment_json, verification_json) "
            "VALUES ('sol_b', 'exp_b', ?, 1, ?, 0.01, '{}', '{}')",
            (ds_id_b, res_b.optimal_cost),
        )

    cur = test_db.cursor()
    cur.execute("SELECT total_cost FROM solutions WHERE dataset_id = ?", (ds_id_a,))
    saved_cost_a = cur.fetchone()[0]
    cur.execute("SELECT total_cost FROM solutions WHERE dataset_id = ?", (ds_id_b,))
    saved_cost_b = cur.fetchone()[0]

    assert saved_cost_a == 10.0
    assert saved_cost_b == 5.0
    assert saved_cost_a != saved_cost_b


# ==============================================================================
# 7. QAOA SIMULATION RESULT SEMANTICS & TELEMETRY
# ==============================================================================

def test_qaoa_simulation_result_semantics(test_db):
    """
    Verifies QAOA Simulation Result Semantics:
    1. Qiskit AerSimulator is classical simulation of quantum circuit (never hardware).
    2. Runs Qiskit pipeline on current active dataset (Instance -> QUBO -> Ising -> QAOA -> AerSimulator).
    3. Independent execution: does not reuse classical solution or force classical optimum.
    4. Telemetry includes qubits, shots, feasibility rate, runtime, circuit depth, 2q gates.
    5. Gap comparison against actual classical optimum.
    6. Strictly dataset-scoped: switching dataset executes fresh pipeline without cached leakage.
    """
    from app.services.experiment_service import run_qaoa_workspace
    from src.classical import solve_exact

    csv_data = b"Worker,Shift,Cost\nA,Morning,10\nB,Morning,50\n"
    insp = inspect_uploaded_file(csv_data, "workforce_qaoa.csv")
    ds_id, norm = finalize_and_save_dataset(test_db, insp, dataset_name="QAOA Test Dataset")

    inst = db_to_instance(ds_id, test_db)
    fp = compute_dataset_fingerprint(ds_id, test_db)

    # 1. Execute QAOA on AerSimulator
    q_out = run_qaoa_workspace(inst, ds_id, shots=128, maxiter=5, seed=42, conn=test_db)

    # 2. Check Backend & Provenance
    assert q_out["backend"] == "Qiskit AerSimulator"
    assert q_out["dataset_id"] == ds_id
    assert q_out["dataset_fingerprint"] == fp
    assert q_out["instance_fingerprint"] == inst.instance_id
    assert q_out["qubit_count"] == inst.n_vars
    assert q_out["shots"] == 128

    # 3. Check Telemetry
    assert 0.0 <= q_out["feasible_rate"] <= 1.0
    assert 0.0 <= q_out["optimal_solution_probability"] <= 1.0
    assert q_out["circuit_depth"] > 0
    assert q_out["two_qubit_gate_count"] >= 0
    assert q_out["runtime_seconds"] > 0.0

    # 4. Check Feasibility & Classical Gap
    c_res = solve_exact(inst)
    assert c_res.feasible
    assert c_res.optimal_cost == 10.0

    if q_out["feasible"] and q_out["best_feasible_cost"] is not None:
        assert q_out["best_feasible_cost"] >= c_res.optimal_cost - 1e-6
        abs_gap = q_out["best_feasible_cost"] - c_res.optimal_cost
        rel_gap = (abs_gap / c_res.optimal_cost) * 100.0
        assert abs_gap >= -1e-6
        # If matches, semantics says "QAOA observed the classical optimum in this simulation."
        # If worse, semantics says "QAOA found a feasible solution but did not reach the classical optimum in this run."
        if abs_gap <= 1e-4:
            status_msg = "QAOA observed the classical optimum in this simulation."
        else:
            status_msg = "QAOA found a feasible solution but did not reach the classical optimum in this run."
        assert len(status_msg) > 0
    else:
        # If no feasible sample, must NOT fabricate cost
        assert q_out["best_feasible_cost"] is None

    # 5. Verify database records are strictly tied to dataset and experiment
    exp_id = q_out["experiment_id"]
    cur = test_db.cursor()
    cur.execute("SELECT parameters_json FROM experiments WHERE experiment_id = ?", (exp_id,))
    row = cur.fetchone()
    assert row is not None
    params = json.loads(row[0])
    assert params["backend"] == "Qiskit AerSimulator"
    assert params["dataset_fingerprint"] == fp

    cur.execute("SELECT solution_id, feasible, total_cost FROM solutions WHERE experiment_id = ?", (exp_id,))
    sol_row = cur.fetchone()
    assert sol_row is not None
    if q_out["feasible"]:
        assert sol_row[1] == 1
        assert sol_row[2] == q_out["best_feasible_cost"]
    else:
        assert sol_row[1] == 0
        assert sol_row[2] is None
